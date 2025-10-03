import os
import mimetypes
import zipfile
import tarfile
import fitz  # PyMuPDF
import docx
import openpyxl
import pptx
import json
import xml.etree.ElementTree as ET
import yaml
from flask import Flask, request, render_template_string, send_file
from werkzeug.utils import secure_filename

app = Flask(__name__)
UPLOAD_FOLDER = "uploads"
EXTRACT_FOLDER = "extracted"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(EXTRACT_FOLDER, exist_ok=True)

# HTMLテンプレート
HTML = """
<!DOCTYPE html>
<html>
<head>
    <title>ファイル展開サイト</title>
</head>
<body>
    <h1>ファイル展開サービス</h1>
    <form action="/upload" method="post" enctype="multipart/form-data">
        <input type="file" name="file" required>
        <input type="submit" value="アップロード">
    </form>
    {% if message %}
        <h2>結果:</h2>
        <pre>{{ message }}</pre>
    {% endif %}
</body>
</html>
"""

def handle_text(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()[:2000]

def handle_pdf(file_path):
    doc = fitz.open(file_path)
    text = ""
    for page in doc:
        text += page.get_text()
    return text[:2000]

def handle_docx(file_path):
    doc = docx.Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])[:2000]

def handle_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(str(row))
    return "\n".join(rows)[:2000]

def handle_pptx(file_path):
    prs = pptx.Presentation(file_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)[:2000]

def handle_zip(file_path):
    with zipfile.ZipFile(file_path, "r") as zip_ref:
        zip_ref.extractall(EXTRACT_FOLDER)
    return "ZIPファイルを展開しました。"

def handle_tar(file_path):
    with tarfile.open(file_path, "r:*") as tar_ref:
        tar_ref.extractall(EXTRACT_FOLDER)
    return "TARファイルを展開しました。"

def handle_json(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return json.dumps(json.load(f), indent=2)[:2000]

def handle_xml(file_path):
    tree = ET.parse(file_path)
    root = tree.getroot()
    return ET.tostring(root, encoding="utf-8").decode("utf-8")[:2000]

def handle_yaml(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return str(data)[:2000]

def handle_binary(file_path):
    with open(file_path, "rb") as f:
        data = f.read(200)
    return f"バイナリデータ(最初の200バイト): {data.hex()}"

# 拡張子ごとの対応
handlers = {
    # テキスト
    ".txt": handle_text,
    ".md": handle_text,
    ".csv": handle_text,
    ".json": handle_json,
    ".xml": handle_xml,
    ".yaml": handle_yaml,
    ".yml": handle_yaml,
    # 画像は中身を展開せずに情報だけ返す
    ".png": handle_binary,
    ".jpg": handle_binary,
    ".jpeg": handle_binary,
    ".gif": handle_binary,
    ".bmp": handle_binary,
    ".tiff": handle_binary,
    ".webp": handle_binary,
    # PDF / Office
    ".pdf": handle_pdf,
    ".docx": handle_docx,
    ".xlsx": handle_xlsx,
    ".pptx": handle_pptx,
    # アーカイブ
    ".zip": handle_zip,
    ".tar": handle_tar,
    ".gz": handle_tar,
    ".bz2": handle_tar,
    # バイナリ系
    ".exe": handle_binary,
    ".dll": handle_binary,
    ".so": handle_binary,
    ".bin": handle_binary,
    ".mp3": handle_binary,
    ".wav": handle_binary,
    ".flac": handle_binary,
    ".ogg": handle_binary,
    ".mp4": handle_binary,
    ".avi": handle_binary,
    ".mov": handle_binary,
    ".mkv": handle_binary
}

@app.route("/", methods=["GET"])
def index():
    return render_template_string(HTML)

@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return render_template_string(HTML, message="ファイルがありません")

    file = request.files["file"]
    if file.filename == "":
        return render_template_string(HTML, message="ファイル名が空です")

    filename = secure_filename(file.filename)
    file_path = os.path.join(UPLOAD_FOLDER, filename)
    file.save(file_path)

    ext = os.path.splitext(filename)[1].lower()
    if ext in handlers:
        try:
            message = handlers[ext](file_path)
        except Exception as e:
            message = f"処理中にエラーが発生しました: {str(e)}"
    else:
        message = "このファイル形式には対応できません"

    return render_template_string(HTML, message=message)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=10000)
